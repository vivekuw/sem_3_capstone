import docx
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls
import os

doc = docx.Document()
section = doc.sections[0]
footer = section.footer
p = footer.paragraphs[0]
p.alignment = docx.enum.text.WD_ALIGN_PARAGRAPH.CENTER
run = p.add_run()
fldChar1 = parse_xml(r'<w:fldChar %s w:fldCharType="begin"/>' % nsdecls('w'))
instrText = parse_xml(r'<w:instrText %s xml:space="preserve"> PAGE </w:instrText>' % nsdecls('w'))
fldChar2 = parse_xml(r'<w:fldChar %s w:fldCharType="separate"/>' % nsdecls('w'))
fldChar3 = parse_xml(r'<w:fldChar %s w:fldCharType="end"/>' % nsdecls('w'))
run._r.append(fldChar1)
run._r.append(instrText)
run._r.append(fldChar2)
run._r.append(fldChar3)

doc.add_paragraph("Hello Page 1")
doc.save("test_page_num.docx")
print("Page num test passed!")
if os.path.exists("test_page_num.docx"):
    os.remove("test_page_num.docx")

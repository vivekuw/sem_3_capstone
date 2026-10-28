import os
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import parse_xml, OxmlElement
from docx.oxml.ns import nsdecls, qn

RUNNING_TITLE = "“Explainable Industrial AI Root-Cause Agent using Temporal and Knowledge Graphs for Continuous Chemical Processes”"

def set_page_borders(section):
    sectPr = section._sectPr
    border_xml = (
        f'<w:pgBorders {nsdecls("w")} w:offsetFrom="page">'
        '<w:top w:val="single" w:sz="12" w:space="24" w:color="000000"/>'
        '<w:left w:val="single" w:sz="12" w:space="24" w:color="000000"/>'
        '<w:bottom w:val="single" w:sz="12" w:space="24" w:color="000000"/>'
        '<w:right w:val="single" w:sz="12" w:space="24" w:color="000000"/>'
        '</w:pgBorders>'
    )
    sectPr.append(parse_xml(border_xml))

def set_cell_background(cell, hex_color):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{hex_color}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=140, right=140):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(
        f'<w:tcMar {nsdecls("w")}>'
        f'<w:top w:w="{top}" w:type="dxa"/>'
        f'<w:bottom w:w="{bottom}" w:type="dxa"/>'
        f'<w:left w:w="{left}" w:type="dxa"/>'
        f'<w:right w:w="{right}" w:type="dxa"/>'
        f'</w:tcMar>'
    )
    tcPr.append(tcMar)

def add_chapter_header(doc, chapter_num, chapter_title):
    # Top Right Chapter indicator
    p_num = doc.add_paragraph()
    p_num.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    p_num.paragraph_format.space_before = Pt(0)
    p_num.paragraph_format.space_after = Pt(4)
    r_num = p_num.add_run(f"Chapter {chapter_num}")
    r_num.font.name = "Times New Roman"
    r_num.font.size = Pt(16)
    r_num.font.bold = True
    r_num.font.color.rgb = RGBColor(0x0F, 0x17, 0x2A)

    # Top line rule
    p_rule = doc.add_paragraph()
    p_rule.paragraph_format.space_before = Pt(0)
    p_rule.paragraph_format.space_after = Pt(14)
    p_rule_pPr = p_rule._p.get_or_add_pPr()
    pBdr = parse_xml(f'<w:pBdr {nsdecls("w")}><w:bottom w:val="single" w:sz="12" w:space="1" w:color="0F172A"/></w:pBdr>')
    p_rule_pPr.append(pBdr)

    # Centered Big Chapter Title
    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_title.paragraph_format.space_before = Pt(6)
    p_title.paragraph_format.space_after = Pt(16)
    r_title = p_title.add_run(chapter_title.upper())
    r_title.font.name = "Times New Roman"
    r_title.font.size = Pt(18)
    r_title.font.bold = True
    r_title.font.color.rgb = RGBColor(0x0F, 0x17, 0x2A)

def add_section_heading(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(14)
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.keep_with_next = True
    run = p.add_run(text)
    run.font.name = "Times New Roman"
    run.font.size = Pt(14)
    run.font.bold = True
    run.font.color.rgb = RGBColor(0x0F, 0x17, 0x2A)
    return p

def add_subsection_heading(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(10)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.keep_with_next = True
    run = p.add_run(text)
    run.font.name = "Times New Roman"
    run.font.size = Pt(12)
    run.font.bold = True
    run.font.color.rgb = RGBColor(0x1E, 0x29, 0x3B)
    return p

def add_paragraph(doc, text, bold_prefix=None, space_after=8, line_spacing=1.25):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.line_spacing = line_spacing
    
    if bold_prefix:
        r_pre = p.add_run(bold_prefix + " ")
        r_pre.font.name = "Times New Roman"
        r_pre.font.size = Pt(12)
        r_pre.font.bold = True
        r_pre.font.color.rgb = RGBColor(0x0F, 0x17, 0x2A)
        
    run = p.add_run(text)
    run.font.name = "Times New Roman"
    run.font.size = Pt(12)
    run.font.color.rgb = RGBColor(0x1E, 0x29, 0x3B)
    return p

def add_bullet(doc, text, bold_prefix=None):
    p = doc.add_paragraph(style='List Bullet')
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(5)
    p.paragraph_format.line_spacing = 1.25
    
    if bold_prefix:
        r_pre = p.add_run(bold_prefix + ": ")
        r_pre.font.name = "Times New Roman"
        r_pre.font.size = Pt(11.5)
        r_pre.font.bold = True
        r_pre.font.color.rgb = RGBColor(0x0F, 0x17, 0x2A)
        
    run = p.add_run(text)
    run.font.name = "Times New Roman"
    run.font.size = Pt(11.5)
    run.font.color.rgb = RGBColor(0x1E, 0x29, 0x3B)
    return p

def add_code_block(doc, code_text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(8)
    p.paragraph_format.left_indent = Inches(0.4)
    p_pPr = p._p.get_or_add_pPr()
    pBdr = parse_xml(f'<w:pBdr {nsdecls("w")}><w:left w:val="single" w:sz="18" w:space="8" w:color="38BDF8"/></w:pBdr>')
    p_pPr.append(pBdr)
    run = p.add_run(code_text)
    run.font.name = "Consolas"
    run.font.size = Pt(9.5)
    run.font.color.rgb = RGBColor(0x0F, 0x17, 0x2A)
    return p

def add_figure(doc, image_path, caption_text, width_inches=6.0):
    if not os.path.exists(image_path):
        print(f"Warning: image {image_path} not found.")
        return
    p_img = doc.add_paragraph()
    p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_img.paragraph_format.space_before = Pt(8)
    p_img.paragraph_format.space_after = Pt(4)
    run_img = p_img.add_run()
    run_img.add_picture(image_path, width=Inches(width_inches))
    
    p_cap = doc.add_paragraph()
    p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_cap.paragraph_format.space_before = Pt(0)
    p_cap.paragraph_format.space_after = Pt(10)
    run_cap = p_cap.add_run(caption_text)
    run_cap.font.name = "Times New Roman"
    run_cap.font.size = Pt(10)
    run_cap.font.bold = True
    run_cap.font.italic = True
    run_cap.font.color.rgb = RGBColor(0x33, 0x41, 0x55)

def add_table_data(doc, headers, data, col_widths, align_left_col0=True):
    table = doc.add_table(rows=len(data) + 1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    
    # Table border xml
    tblPr = table._tbl.tblPr
    tblBorders = parse_xml(
        f'<w:tblBorders {nsdecls("w")}>'
        '<w:top w:val="single" w:sz="8" w:space="0" w:color="334155"/>'
        '<w:left w:val="single" w:sz="8" w:space="0" w:color="334155"/>'
        '<w:bottom w:val="single" w:sz="8" w:space="0" w:color="334155"/>'
        '<w:right w:val="single" w:sz="8" w:space="0" w:color="334155"/>'
        '<w:insideH w:val="single" w:sz="4" w:space="0" w:color="CBD5E1"/>'
        '<w:insideV w:val="single" w:sz="4" w:space="0" w:color="CBD5E1"/>'
        '</w:tblBorders>'
    )
    tblPr.append(tblBorders)
    
    # Header Row
    hdr_cells = table.rows[0].cells
    for i, h_text in enumerate(headers):
        hdr_cells[i].text = h_text
        set_cell_background(hdr_cells[i], "1E3A8A") # Navy Blue
        set_cell_margins(hdr_cells[i], top=120, bottom=120, left=140, right=140)
        p = hdr_cells[i].paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(0)
        for run in p.runs:
            run.font.name = "Times New Roman"
            run.font.size = Pt(10)
            run.font.bold = True
            run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
            
    # Set header repeat on every page
    trPr = table.rows[0]._tr.get_or_add_trPr()
    trPr.append(parse_xml(f'<w:tblHeader {nsdecls("w")}/>'))

    # Data Rows
    for row_idx, row_values in enumerate(data):
        row_cells = table.rows[row_idx + 1].cells
        bg_col = "F8FAFC" if row_idx % 2 == 1 else "FFFFFF"
        for col_idx, cell_value in enumerate(row_values):
            row_cells[col_idx].text = str(cell_value)
            set_cell_background(row_cells[col_idx], bg_col)
            set_cell_margins(row_cells[col_idx], top=80, bottom=80, left=120, right=120)
            p = row_cells[col_idx].paragraphs[0]
            if col_idx == 0 and align_left_col0:
                p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            else:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER if col_idx in [1, 2, 4] else WD_ALIGN_PARAGRAPH.LEFT
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.1
            for run in p.runs:
                run.font.name = "Times New Roman"
                run.font.size = Pt(9.5)
                run.font.color.rgb = RGBColor(0x1E, 0x29, 0x3B)
                
    # Column widths
    for row in table.rows:
        for idx, width in enumerate(col_widths):
            row.cells[idx].width = Inches(width)
            
    p_space = doc.add_paragraph()
    p_space.paragraph_format.space_before = Pt(0)
    p_space.paragraph_format.space_after = Pt(8)

def setup_footer(section, is_preliminary=False, roman_page_str=None):
    footer = section.footer
    p = footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(2)
    
    if is_preliminary:
        if roman_page_str:
            r = p.add_run(roman_page_str)
            r.font.name = "Times New Roman"
            r.font.size = Pt(10)
            r.font.color.rgb = RGBColor(0x47, 0x55, 0x69)
    else:
        # Running page number
        r_num = p.add_run()
        fldChar1 = parse_xml(r'<w:fldChar %s w:fldCharType="begin"/>' % nsdecls('w'))
        instrText = parse_xml(r'<w:instrText %s xml:space="preserve"> PAGE </w:instrText>' % nsdecls('w'))
        fldChar2 = parse_xml(r'<w:fldChar %s w:fldCharType="separate"/>' % nsdecls('w'))
        fldChar3 = parse_xml(r'<w:fldChar %s w:fldCharType="end"/>' % nsdecls('w'))
        r_num._r.append(fldChar1)
        r_num._r.append(instrText)
        r_num._r.append(fldChar2)
        r_num._r.append(fldChar3)
        r_num.font.name = "Times New Roman"
        r_num.font.size = Pt(10)
        r_num.font.bold = True
        r_num.font.color.rgb = RGBColor(0x0F, 0x17, 0x2A)
        
        # Quote line below
        p_quote = footer.add_paragraph()
        p_quote.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_quote.paragraph_format.space_before = Pt(0)
        p_quote.paragraph_format.space_after = Pt(0)
        r_quote = p_quote.add_run(RUNNING_TITLE)
        r_quote.font.name = "Times New Roman"
        r_quote.font.size = Pt(8.5)
        r_quote.font.italic = True
        r_quote.font.color.rgb = RGBColor(0x64, 0x74, 0x8B)
